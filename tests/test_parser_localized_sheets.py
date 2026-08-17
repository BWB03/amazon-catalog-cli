"""Tests for localized Amazon CLR worksheet discovery."""

import openpyxl
import pytest

from catalog.core.parser import CLRParser


LOCALIZED_TEMPLATE_SHEETS = (
    ("DE", "Vorlage"),
    ("FR", "Modèle"),
    ("IT", "Modello"),
    ("ES", "Plantilla"),
    ("SE", "Mall"),
    ("PL", "Szablon"),
    ("BE", "Sjabloon"),
    ("NL", "Sjabloon"),
)


def _write_clr(path, sheet_name, marketplace="US"):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.cell(row=1, column=1, value=f"Country:{marketplace}")

    headers = ["Status", "Title", "SKU", "Product Type"]
    field_ids = ["status#1.value", "item_name#1.value", "item_sku", "product_type"]
    values = ["Active", "Test product", "SKU-1", "TEST_PRODUCT"]
    for column, (header, field_id, value) in enumerate(
        zip(headers, field_ids, values), start=1
    ):
        sheet.cell(row=4, column=column, value=header)
        sheet.cell(row=5, column=column, value=field_id)
        sheet.cell(row=7, column=column, value=value)

    workbook.save(path)


@pytest.mark.parametrize(("marketplace", "sheet_name"), LOCALIZED_TEMPLATE_SHEETS)
def test_parser_supports_localized_template_sheet_names(tmp_path, marketplace, sheet_name):
    clr_path = tmp_path / f"{marketplace.lower()}-catalog.xlsx"
    _write_clr(clr_path, sheet_name, marketplace)

    parser = CLRParser(str(clr_path))

    assert parser.template_sheet.title == sheet_name
    assert parser.get_marketplace() == marketplace
    assert parser.headers["SKU"] == 3
    assert parser.get_listings()[0].sku == "SKU-1"


def test_parser_matches_localized_sheet_name_case_and_accents(tmp_path):
    clr_path = tmp_path / "fr-catalog.xlsx"
    _write_clr(clr_path, "modele", "FR")

    parser = CLRParser(str(clr_path))

    assert parser.template_sheet.title == "modele"


@pytest.mark.parametrize("marketplace", ("SE", "PL", "BE", "NL"))
def test_parser_recognizes_new_marketplace_codes_without_country_prefix(tmp_path, marketplace):
    clr_path = tmp_path / f"{marketplace.lower()}-catalog.xlsx"
    _write_clr(clr_path, "Template", marketplace)
    workbook = openpyxl.load_workbook(clr_path)
    workbook["Template"].cell(row=1, column=1, value=marketplace)
    workbook.save(clr_path)

    parser = CLRParser(str(clr_path))

    assert parser.get_marketplace() == marketplace


def test_parser_uses_structural_fallback_for_unknown_localized_name(tmp_path):
    clr_path = tmp_path / "future-marketplace.xlsx"
    _write_clr(clr_path, "Localized Template Name", "US")

    parser = CLRParser(str(clr_path))

    assert parser.template_sheet.title == "Localized Template Name"


def test_parser_does_not_fall_back_to_an_arbitrary_worksheet(tmp_path):
    clr_path = tmp_path / "not-a-clr.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Instructions"
    workbook.save(clr_path)

    with pytest.raises(ValueError, match="Could not find the Amazon CLR template worksheet"):
        CLRParser(str(clr_path))
