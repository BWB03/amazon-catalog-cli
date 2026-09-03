"""Copy-aware CLR parsing and missing-attribute behavior."""

from openpyxl import Workbook

from catalog.core.parser import CLRParser
from catalog.queries.missing_attributes import MissingAttributesQuery


HEADERS = [
    "Status",
    "Title",
    "SKU",
    "Product Type",
    "Listing Action",
    "Parentage",
    "Parent SKU",
    "Brand",
    "Product Id Type",
    "Product Id",
]


def _write_clr(tmp_path, rows):
    path = tmp_path / "copy-aware.xlsx"
    workbook = Workbook()
    template = workbook.active
    template.title = "Template"
    template.cell(1, 1, "Country:US")
    for column, header in enumerate(HEADERS, start=1):
        template.cell(4, column, header)
        template.cell(5, column, header)
    template.cell(6, 3, "EXAMPLE-SKU")
    for row_number, values in enumerate(rows, start=7):
        for column, value in enumerate(values, start=1):
            template.cell(row_number, column, value)

    definitions = workbook.create_sheet("Data Definitions")
    definitions.append(["Group Name", "Field Name", "Required?"])
    definitions.append([
        "Reference Group - This group contains reference-only attributes.",
        "Status",
        "Required",
    ])
    definitions.append(["Offer", "Brand", "Required"])
    workbook.save(path)
    return path


def test_reference_only_required_field_is_not_a_seller_fix(tmp_path):
    path = _write_clr(tmp_path, [["", "Product", "SKU-1", "FOOD", "", "", "", "", "ASIN", "B000000001"]])
    parser = CLRParser(str(path))

    assert parser.get_required_fields() == ["Brand"]


def test_same_title_distinct_skus_are_preserved(tmp_path):
    path = _write_clr(tmp_path, [
        ["Active", "Same title", "SKU-A", "FOOD", "", "", "", "Brand", "ASIN", "B000000001"],
        ["Active", "Same title", "SKU-B", "FOOD", "", "", "", "Brand", "ASIN", "B000000002"],
    ])
    parser = CLRParser(str(path))

    listings = parser.get_listings()
    metadata = parser.get_listing_filter_metadata()["record_equivalence"]

    assert [listing.sku for listing in listings] == ["SKU-A", "SKU-B"]
    assert metadata["same_title_candidate_clusters"] == 1
    assert metadata["title_only_rows_excluded"] == 0


def test_exact_sku_duplicate_uses_populated_copy_without_false_blank(tmp_path):
    path = _write_clr(tmp_path, [
        ["Active", "Product", "SKU-DUP", "FOOD", "", "", "", "Brand", "ASIN", "B000000003"],
        ["Active", "", "SKU-DUP", "FOOD", "", "", "", "", "ASIN", "B000000003"],
    ])
    parser = CLRParser(str(path))
    listings = parser.get_listings()

    issues = MissingAttributesQuery().execute(listings, parser)
    metadata = parser.get_listing_filter_metadata()["record_equivalence"]

    assert issues == []
    assert len(listings) == 1
    assert metadata["exact_sku_rows_consolidated"] == 1
    assert metadata["exact_sku_duplicate_clusters"][0]["sku"] == "SKU-DUP"


def test_shared_asin_blank_is_review_metadata_not_confirmed_gap(tmp_path):
    path = _write_clr(tmp_path, [
        ["Active", "Product", "OFFER-A", "FOOD", "", "", "", "Brand", "ASIN", "B000000004"],
        ["Active", "Product", "OFFER-B", "FOOD", "", "", "", "", "ASIN", "B000000004"],
    ])
    parser = CLRParser(str(path))
    listings = parser.get_listings()

    issues = MissingAttributesQuery().execute(listings, parser)
    metadata = parser.get_listing_filter_metadata()

    assert issues == []
    assert len(listings) == 2
    assert metadata["record_equivalence"]["shared_asin_review_clusters"][0]["skus"] == ["OFFER-A", "OFFER-B"]
    assert metadata["copy_aware_blank_reviews"][0]["match_type"] == "shared_asin_review"


def test_removed_row_does_not_create_active_content_gap(tmp_path):
    path = _write_clr(tmp_path, [["Removed", "Product", "OLD-SKU", "FOOD", "Delete", "", "", "", "ASIN", "B000000005"]])
    parser = CLRParser(str(path))

    issues = MissingAttributesQuery().execute(parser.get_listings(), parser)

    assert issues == []
