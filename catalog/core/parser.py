"""
CLR Parser Module
Extracts and normalizes data from Amazon Category Listing Reports
"""

import openpyxl
from typing import Optional, Dict, List
from dataclasses import dataclass

# Suppress stderr output (set True for MCP stdio transport)
_quiet = False


@dataclass
class Listing:
    """Normalized listing data"""
    row_number: int
    sku: str
    product_type: str
    item_type: str
    title: str
    brand: str
    parentage: str
    parent_sku: str
    status: str
    bullet_points: List[str]
    all_fields: Dict[str, any]  # Full row data


class CLRParser:
    """Parse Amazon Category Listing Reports"""
    
    # Standard CLR row structure
    ROW_SETTINGS = 1
    ROW_INSTRUCTIONS = 2
    ROW_GROUP_HEADERS = 3
    ROW_COL_HEADERS = 4
    ROW_FIELD_IDS = 5
    ROW_EXAMPLE = 6
    ROW_DATA_START = 7
    PRODUCT_ID_FIELD_NAMES = {"product id", "product id type"}
    
    def __init__(self, clr_file_path: str):
        """Load and parse CLR file"""
        self.file_path = clr_file_path
        self.workbook = openpyxl.load_workbook(clr_file_path, data_only=True, read_only=True)
        self._listing_filter_metadata = {}
        self._unfiltered_listings = []
        self._copy_aware_blank_reviews = []
        
        # Load sheets
        self.template_sheet = self.workbook['Template']
        
        # Parse structure
        self.headers = self._extract_headers()
        self._field_id_to_display = self._build_field_id_map()
        self.field_definitions = self._extract_field_definitions()
        self.marketplace = self._extract_marketplace()
        
    def _extract_marketplace(self) -> str:
        """
        Extract marketplace from settings row (row 1).
        Amazon CLRs contain marketplace info in row 1, typically formatted as:
        "Country:US" or similar.
        
        Returns:
            str: Marketplace code (e.g., 'US', 'CA', 'UK', 'DE') or 'US' as default
        """
        try:
            settings_row = self.template_sheet[self.ROW_SETTINGS]
            
            # Scan first 10 cells for country/marketplace info
            for cell in list(settings_row)[:10]:
                if cell.value:
                    value = str(cell.value).strip()
                    
                    # Look for "Country:XX" pattern
                    if 'country:' in value.lower():
                        parts = value.split(':')
                        if len(parts) == 2:
                            return parts[1].strip().upper()
                    
                    # Look for direct marketplace codes
                    value_upper = value.upper()
                    known_marketplaces = ['US', 'CA', 'UK', 'DE', 'FR', 'IT', 'ES', 'JP', 'AU', 'IN', 'MX', 'BR']
                    if value_upper in known_marketplaces:
                        return value_upper
            
            # Default to US if not found
            return 'US'
            
        except Exception:
            # If any error, default to US
            return 'US'
    
    def get_marketplace(self) -> str:
        """Get the marketplace code for this CLR"""
        return self.marketplace
    
    def is_us_marketplace(self) -> bool:
        """Check if this is a US marketplace CLR"""
        return self.marketplace == 'US'
        
    def _extract_headers(self) -> Dict[str, int]:
        """Extract column headers and their positions"""
        headers = {}
        header_row = self.template_sheet[self.ROW_COL_HEADERS]
        
        for idx, cell in enumerate(header_row, start=1):
            if cell.value:
                headers[str(cell.value).strip()] = idx
        
        return headers
    
    def _build_field_id_map(self) -> Dict[str, str]:
        """Build a mapping from Row 5 field IDs to Row 4 display names.

        CLR files have:
          Row 4: Display names (e.g., 'Item Name', 'Brand Name')
          Row 5: Field IDs (e.g., 'item_name[marketplace_id=...]#1.value')

        Data Definitions uses field IDs, but all_fields uses display names.
        This map lets us translate between them.
        """
        field_id_map = {}
        try:
            header_row = self.template_sheet[self.ROW_COL_HEADERS]
            field_id_row = self.template_sheet[self.ROW_FIELD_IDS]

            for header_cell, field_id_cell in zip(header_row, field_id_row):
                if header_cell.value and field_id_cell.value:
                    field_id = str(field_id_cell.value).strip()
                    display_name = str(header_cell.value).strip()
                    field_id_map[field_id] = display_name
        except Exception:
            pass
        return field_id_map

    def _extract_field_definitions(self) -> Dict[str, Dict]:
        """Extract field definitions from Data Definitions sheet"""
        definitions = {}
        
        try:
            dd_sheet = self.workbook['Data Definitions']
            
            # Find header row (usually row 1, but scan first 5)
            header_row_idx = None
            for i in range(1, 6):
                row = dd_sheet[i]
                if any(cell.value and 'field name' in str(cell.value).lower() for cell in row):
                    header_row_idx = i
                    break
            
            if not header_row_idx:
                return definitions
            
            # Extract headers
            dd_headers = {}
            for idx, cell in enumerate(dd_sheet[header_row_idx], start=1):
                if cell.value:
                    dd_headers[str(cell.value).strip().lower()] = idx
            
            # Extract definitions. Group headings are often carried only on the
            # first row of a section, so preserve the most recent nonblank value.
            current_group = ""
            for row in dd_sheet.iter_rows(min_row=header_row_idx + 1):
                field_name_idx = dd_headers.get('field name')
                required_idx = dd_headers.get('required?')
                group_idx = dd_headers.get('group name') or dd_headers.get('group')
                
                if not field_name_idx:
                    continue
                
                group = row[group_idx - 1].value if group_idx else None
                if group:
                    current_group = str(group).strip()

                field_name = row[field_name_idx - 1].value
                if not field_name:
                    continue
                
                required = row[required_idx - 1].value if required_idx else None
                
                definitions[str(field_name).strip()] = {
                    'required': str(required).strip().lower() if required else '',
                    'field_name': str(field_name).strip(),
                    'group': current_group,
                }
        
        except KeyError:
            # No Data Definitions sheet
            pass
        
        return definitions
    
    def _resolve_field_name(self, field_id: str) -> Optional[str]:
        """Resolve a field ID to its display name.

        If the field ID matches a display name directly, return it.
        Otherwise look it up in the field_id_to_display map.
        Returns None if the field doesn't exist in this CLR's columns.
        """
        # Direct match (field ID is already a display name)
        if field_id in self.headers:
            return field_id
        # Map from field ID to display name
        display = self._field_id_to_display.get(field_id)
        if display and display in self.headers:
            return display
        return None

    def get_required_fields(self) -> List[str]:
        """Get list of required field display names that exist in this CLR."""
        fields = []
        for field_name, definition in self.field_definitions.items():
            if (
                definition['required'] == 'required'
                and not self._is_reference_only_definition(definition)
            ):
                resolved = self._resolve_field_name(field_name)
                if resolved:
                    fields.append(resolved)
        return fields

    def get_conditional_fields(self) -> List[str]:
        """Get list of conditionally required field display names that exist in this CLR."""
        fields = []
        for field_name, definition in self.field_definitions.items():
            if (
                'conditional' in definition['required'].lower()
                and not self._is_reference_only_definition(definition)
            ):
                resolved = self._resolve_field_name(field_name)
                if resolved:
                    fields.append(resolved)
        return fields

    @staticmethod
    def _is_reference_only_definition(definition: Dict) -> bool:
        """Return True when Amazon places a field in a reference-only group."""
        group = str(definition.get('group') or '').strip().lower()
        return 'reference-only' in group or 'reference only' in group

    @classmethod
    def is_product_identifier_field(cls, field_name: str) -> bool:
        """Return True for CLR Product Id / Product Id Type display fields."""
        return field_name.strip().lower() in cls.PRODUCT_ID_FIELD_NAMES

    def is_virtual_bundle_listing(self, listing: Listing) -> bool:
        """
        Identify Amazon virtual bundles from CLR identifier cells.

        Virtual bundles do not carry Product Id or Product Id Type values in the
        report, while regular listings use UPC, ASIN, or GTIN Exempt identifiers.
        """
        identifier_fields = [
            field
            for field in listing.all_fields
            if self.is_product_identifier_field(field)
        ]
        normalized_identifier_fields = {
            field.strip().lower()
            for field in identifier_fields
        }
        if normalized_identifier_fields != self.PRODUCT_ID_FIELD_NAMES:
            return False

        return all(
            not str(listing.all_fields.get(field) or "").strip()
            for field in identifier_fields
        )
    
    def get_listings(self, skip_parents: bool = True, skip_examples: bool = True, skip_fbm_duplicates: bool = True) -> List[Listing]:
        """
        Extract all listings from CLR
        
        Args:
            skip_parents: Skip parent SKUs (variations)
            skip_examples: Skip example/dummy rows
        
        Returns:
            List of normalized Listing objects
        """
        listings = []
        
        # Get column indices
        col_sku = self.headers.get('SKU', 3)
        col_status = self.headers.get('Status', 1)
        col_title = self.headers.get('Title', 2)
        col_product_type = self.headers.get('Product Type', 4)
        col_item_type = self.headers.get('Item Type Keyword', 13)
        col_brand = self.headers.get('Brand', 10)
        col_parentage = self.headers.get('Parentage', 6)
        col_parent_sku = self.headers.get('Parent SKU', 7)
        
        # Bullet point columns — handle both numbered (Bullet Point 1-5) and single (Bullet Point)
        bullet_cols = []
        for i in range(1, 6):
            col_name = f'Bullet Point {i}'
            if col_name in self.headers:
                bullet_cols.append(self.headers[col_name])
        # Fallback: single "Bullet Point" column (some CLR formats)
        if not bullet_cols and 'Bullet Point' in self.headers:
            bullet_cols.append(self.headers['Bullet Point'])
        
        # Iterate through data rows
        for row_idx, row in enumerate(self.template_sheet.iter_rows(min_row=self.ROW_DATA_START), 
                                      start=self.ROW_DATA_START):
            # Extract basic fields
            sku = self._get_cell_value(row, col_sku)
            
            if not sku:
                continue
            
            # Skip examples
            if skip_examples and sku.upper() in ['ABC123', 'EXAMPLE', 'TEST']:
                continue
            
            status = self._get_cell_value(row, col_status)
            parentage = self._get_cell_value(row, col_parentage)
            
            # Skip parents if requested
            if skip_parents and parentage and 'parent' in parentage.lower():
                continue
            
            # Extract all fields
            all_fields = {}
            for header_name, col_idx in self.headers.items():
                all_fields[header_name] = self._get_cell_value(row, col_idx)
            
            # Extract bullet points
            bullets = []
            for bullet_col in bullet_cols:
                bullet_text = self._get_cell_value(row, bullet_col)
                bullets.append(bullet_text if bullet_text else "")
            
            # Create listing object
            listing = Listing(
                row_number=row_idx,
                sku=sku,
                product_type=self._get_cell_value(row, col_product_type) or "",
                item_type=self._get_cell_value(row, col_item_type) or "",
                title=self._get_cell_value(row, col_title) or "",
                brand=self._get_cell_value(row, col_brand) or "",
                parentage=parentage or "",
                parent_sku=self._get_cell_value(row, col_parent_sku) or "",
                status=status or "",
                bullet_points=bullets,
                all_fields=all_fields
            )
            
            listings.append(listing)
        
        self._unfiltered_listings = list(listings)
        self._copy_aware_blank_reviews = []
        self._listing_filter_metadata = self._build_record_equivalence_metadata(listings)

        # Preserve distinct SKUs even when their titles match. When the legacy
        # exclusion switch is enabled, consolidate exact-SKU duplicate records.
        if skip_fbm_duplicates:
            listings = self._filter_fbm_duplicates(listings)
        
        return listings

    def get_listing_filter_metadata(self) -> Dict:
        """Return metadata about row/listing exclusions applied during parsing."""
        metadata = dict(self._listing_filter_metadata)
        metadata['copy_aware_blank_reviews'] = list(self._copy_aware_blank_reviews)
        return metadata

    @staticmethod
    def _normalized(value: Optional[str]) -> str:
        return str(value or '').strip().casefold()

    @staticmethod
    def _get_named_field(listing: Listing, *names: str) -> str:
        wanted = {name.casefold() for name in names}
        for field_name, value in listing.all_fields.items():
            if str(field_name).strip().casefold() in wanted:
                return str(value or '').strip()
        return ''

    def _explicit_asin(self, listing: Listing) -> str:
        identifier_type = self._get_named_field(listing, 'Product Id Type')
        identifier = self._get_named_field(listing, 'Product Id')
        if identifier_type.casefold() != 'asin':
            return ''
        normalized = identifier.strip().upper()
        if len(normalized) == 10 and normalized.isalnum():
            return normalized
        return ''

    @staticmethod
    def _conflicting_fields(group: List[Listing]) -> List[str]:
        fields = set().union(*(listing.all_fields.keys() for listing in group))
        conflicts = []
        for field in fields:
            values = {
                str(listing.all_fields.get(field) or '').strip()
                for listing in group
                if str(listing.all_fields.get(field) or '').strip()
            }
            if len(values) > 1:
                conflicts.append(field)
        return sorted(conflicts)

    def _build_record_equivalence_metadata(self, listings: List[Listing]) -> Dict:
        sku_groups: Dict[str, List[Listing]] = {}
        asin_groups: Dict[str, List[Listing]] = {}
        title_groups: Dict[str, List[Listing]] = {}

        for listing in listings:
            sku_groups.setdefault(self._normalized(listing.sku), []).append(listing)
            asin = self._explicit_asin(listing)
            if asin:
                asin_groups.setdefault(asin, []).append(listing)
            title = self._normalized(listing.title)
            if title:
                title_groups.setdefault(title, []).append(listing)

        duplicate_clusters = [
            {
                'sku': group[0].sku,
                'rows': [listing.row_number for listing in group],
                'conflicting_fields': self._conflicting_fields(group),
            }
            for group in sku_groups.values()
            if len(group) > 1
        ]
        shared_asin_clusters = [
            {
                'asin': asin,
                'skus': sorted({listing.sku for listing in group}),
                'rows': [listing.row_number for listing in group],
                'conflicting_fields': self._conflicting_fields(group),
            }
            for asin, group in asin_groups.items()
            if len({self._normalized(listing.sku) for listing in group}) > 1
        ]
        same_title_candidates = sum(
            1
            for group in title_groups.values()
            if len({self._normalized(listing.sku) for listing in group}) > 1
        )

        return {
            'record_equivalence': {
                'raw_listing_count': len(listings),
                'exact_sku_duplicate_clusters': duplicate_clusters,
                'shared_asin_review_clusters': shared_asin_clusters,
                'same_title_candidate_clusters': same_title_candidates,
                'title_only_rows_excluded': 0,
                'strategy': (
                    'Consolidate exact-SKU duplicate records only. Keep distinct '
                    'SKUs separate; shared ASINs are review clusters and matching '
                    'titles alone never establish equivalence.'
                ),
            },
            'fbm_duplicate_exclusion': {
                'enabled': False,
                'excluded_count': 0,
                'excluded_skus_sample': [],
                'reason': (
                    'Deprecated title-only FBA/FBM filtering is disabled because '
                    'matching titles do not prove record equivalence.'
                ),
                'strategy': 'No title-only exclusions.',
            },
        }
    
    def _filter_fbm_duplicates(self, listings: List[Listing]) -> List[Listing]:
        """
        Consolidate exact-SKU duplicate rows for backwards compatibility.

        The former implementation excluded distinct SKUs based on matching title.
        Titles are not stable identity keys, so distinct SKUs are always preserved.
        """
        grouped: Dict[str, List[Listing]] = {}
        group_order = []
        for listing in listings:
            key = self._normalized(listing.sku)
            if key not in grouped:
                group_order.append(key)
            grouped.setdefault(key, []).append(listing)

        filtered = []
        skipped_rows = []
        for key in group_order:
            group = grouped[key]
            representative = max(
                group,
                key=lambda listing: sum(
                    bool(str(value or '').strip())
                    for value in listing.all_fields.values()
                ),
            )
            filtered.append(representative)
            skipped_rows.extend(
                listing.row_number for listing in group if listing is not representative
            )

        record_metadata = self._listing_filter_metadata['record_equivalence']
        record_metadata['analysis_listing_count'] = len(filtered)
        record_metadata['exact_sku_rows_consolidated'] = len(skipped_rows)
        record_metadata['consolidated_row_numbers'] = skipped_rows[:25]
        return filtered

    def is_removed_or_delete_listing(self, listing: Listing) -> bool:
        """Return True for rows that should not receive active content-gap findings."""
        status = self._normalized(listing.status)
        action = self._normalized(
            self._get_named_field(listing, 'Listing Action', 'Update Delete')
        )
        return status in {'removed', 'deleted'} or action in {'delete', 'removed'}

    def alternate_record_with_value(
        self, listing: Listing, field: str
    ) -> Optional[Dict]:
        """Describe a stronger record containing a value missing on this row."""
        candidates = getattr(self, '_unfiltered_listings', []) or []
        listing_sku = self._normalized(listing.sku)
        listing_asin = self._explicit_asin(listing)

        exact_sku_matches = [
            candidate
            for candidate in candidates
            if candidate.row_number != listing.row_number
            and self._normalized(candidate.sku) == listing_sku
            and str(candidate.all_fields.get(field) or '').strip()
        ]
        if exact_sku_matches:
            return {
                'match_type': 'exact_sku_duplicate',
                'rows': [candidate.row_number for candidate in exact_sku_matches],
                'skus': sorted({candidate.sku for candidate in exact_sku_matches}),
            }

        if listing_asin:
            shared_asin_matches = [
                candidate
                for candidate in candidates
                if candidate.row_number != listing.row_number
                and self._normalized(candidate.sku) != listing_sku
                and self._explicit_asin(candidate) == listing_asin
                and str(candidate.all_fields.get(field) or '').strip()
            ]
            if shared_asin_matches:
                return {
                    'match_type': 'shared_asin_review',
                    'asin': listing_asin,
                    'rows': [candidate.row_number for candidate in shared_asin_matches],
                    'skus': sorted({candidate.sku for candidate in shared_asin_matches}),
                }
        return None

    def record_copy_aware_blank_review(
        self, listing: Listing, field: str, alternate: Dict
    ) -> None:
        """Expose suppressed blanks as review metadata instead of hiding them."""
        review = {
            'row': listing.row_number,
            'sku': listing.sku,
            'field': field,
            **alternate,
        }
        if review not in self._copy_aware_blank_reviews:
            self._copy_aware_blank_reviews.append(review)
    
    def get_product_types(self) -> List[str]:
        """Get unique product types in catalog"""
        product_types = set()
        listings = self.get_listings()
        
        for listing in listings:
            if listing.product_type:
                product_types.add(listing.product_type)
        
        return sorted(list(product_types))
    
    def _get_cell_value(self, row, col_idx: int) -> Optional[str]:
        """Safely get cell value from row"""
        try:
            if col_idx <= 0 or col_idx > len(row):
                return None
            
            cell = row[col_idx - 1]
            if cell.value is None:
                return None
            
            return str(cell.value).strip()
        except (IndexError, AttributeError):
            return None
    
    def __del__(self):
        """Clean up workbook"""
        if hasattr(self, 'workbook'):
            self.workbook.close()
